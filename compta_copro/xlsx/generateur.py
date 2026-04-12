"""Génération d'un classeur Excel (.xlsx) de synthèse.

Le classeur produit contient :
- un onglet par année (totaux par type de charge + ligne TOTAL),
- un onglet ``Comparaison`` (N vs N-1),
- un onglet ``Synthese`` (top hausses en euros et en pourcentage).
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill  # pyright: ignore[reportMissingModuleSource]

from compta_copro.analyse.agregation import comparaison_n_n1, top_hausses

_FORMAT_EURO = '#,##0.00\\ "€"'
_FORMAT_PCT = "0.00%"


def generer_classeur_excel(
    df: pd.DataFrame,
    chemin_sortie: Path,
    annee_cible: int | None = None,
) -> Path:
    """Générer un classeur Excel de synthèse à partir des données normalisées.

    Parameters
    ----------
    df:
        DataFrame normalisé contenant au minimum les colonnes de schéma interne.
    chemin_sortie:
        Chemin du fichier ``.xlsx`` à produire.
    annee_cible:
        Si renseigné, limite l'export à l'onglet de cette année et aux analyses
        de comparaison associées (N versus N-1).

    Returns
    -------
    Path
        Chemin du classeur écrit sur disque.
    """
    chemin_sortie.parent.mkdir(parents=True, exist_ok=True)

    if annee_cible is not None and annee_cible not in set(df["annee"].astype(int).unique()):
        raise ValueError(f"L'annee cible {annee_cible} est absente des donnees")

    comp = comparaison_n_n1(df)
    if annee_cible is not None:
        comp = comp[comp["annee_n"] == annee_cible].reset_index(drop=True)

    with pd.ExcelWriter(chemin_sortie, engine="openpyxl") as writer:
        _ecrire_onglets_par_annee(df, writer, annee_cible=annee_cible)
        _ecrire_onglet_comparaison(comp, writer)
        _ecrire_onglet_synthese(comp, writer)

        classeur = writer.book
        for feuille in classeur.worksheets:
            _formater_entete(feuille)
            _activer_filtres_et_gel(feuille)

    return chemin_sortie


def _ecrire_onglets_par_annee(
    df: pd.DataFrame,
    writer: pd.ExcelWriter,
    annee_cible: int | None = None,
) -> None:
    """Écrire un onglet par année avec les agrégations par type de charge."""
    annees = [annee_cible] if annee_cible is not None else sorted(df["annee"].unique())
    for annee in annees:
        df_annee = df[df["annee"] == annee]
        agg = (
            df_annee.groupby("type_charge", sort=True)
            .agg(
                total_a_repartir=("a_repartir", "sum"),
                total_tva=("tva", "sum"),
                total_recuperable=("recuperable", "sum"),
                nb_lignes=("libelle", "count"),
            )
            .reset_index()
        )
        agg.columns = ["Type de charge", "Total a repartir (EUR)", "TVA (EUR)", "Recuperable (EUR)", "Nb lignes"]

        total = {
            "Type de charge": "TOTAL",
            "Total a repartir (EUR)": agg["Total a repartir (EUR)"].sum(),
            "TVA (EUR)": agg["TVA (EUR)"].sum(),
            "Recuperable (EUR)": agg["Recuperable (EUR)"].sum(),
            "Nb lignes": agg["Nb lignes"].sum(),
        }
        agg = pd.concat([agg, pd.DataFrame([total])], ignore_index=True)
        nom_feuille = str(annee)[:31]
        agg.to_excel(writer, sheet_name=nom_feuille, index=False)

        feuille = writer.book[nom_feuille]
        _ajuster_largeurs(feuille)
        _formater_colonnes_montant(feuille, [2, 3, 4], _FORMAT_EURO)


def _ecrire_onglet_comparaison(comp: pd.DataFrame, writer: pd.ExcelWriter) -> None:
    """Écrire l'onglet de comparaison N vs N-1."""
    if comp.empty:
        pd.DataFrame({"Message": ["Aucune donnee de comparaison disponible."]}).to_excel(
            writer,
            sheet_name="Comparaison",
            index=False,
        )
        return

    df_affichage = comp.rename(
        columns={
            "type_charge": "Type de charge",
            "annee_n": "Annee N",
            "annee_n1": "Annee N-1",
            "total_n": "Total N (EUR)",
            "total_n1": "Total N-1 (EUR)",
            "delta": "Delta (EUR)",
            "delta_pct": "Delta (%)",
        }
    )
    if "cle_repartition" in df_affichage.columns:
        df_affichage = df_affichage.rename(columns={"cle_repartition": "Cle de repartition"})

    df_affichage.to_excel(writer, sheet_name="Comparaison", index=False)
    feuille = writer.book["Comparaison"]
    _ajuster_largeurs(feuille)

    colonnes = [cell.value for cell in feuille[1]]
    indices_eur = [idx + 1 for idx, nom in enumerate(colonnes) if nom and "(EUR)" in str(nom)]
    _formater_colonnes_montant(feuille, indices_eur, _FORMAT_EURO)

    if "Delta (%)" in colonnes:
        idx_pct = colonnes.index("Delta (%)") + 1
        _formater_colonne_pourcentage(feuille, idx_pct)


def _ecrire_onglet_synthese(comp: pd.DataFrame, writer: pd.ExcelWriter) -> None:
    """Écrire l'onglet synthèse avec les tops des hausses."""
    if comp.empty:
        pd.DataFrame({"Message": ["Aucune donnee de comparaison disponible."]}).to_excel(
            writer,
            sheet_name="Synthese",
            index=False,
        )
        return

    top_eur = top_hausses(comp, colonne_tri="delta", n=20)
    top_pct = top_hausses(comp, colonne_tri="delta_pct", n=20)

    top_eur_affichage = top_eur.rename(
        columns={
            "type_charge": "Type de charge",
            "annee_n": "Annee N",
            "annee_n1": "Annee N-1",
            "total_n": "Total N (EUR)",
            "total_n1": "Total N-1 (EUR)",
            "delta": "Delta (EUR)",
            "delta_pct": "Delta (%)",
        }
    )
    top_pct_affichage = top_pct.rename(
        columns={
            "type_charge": "Type de charge",
            "annee_n": "Annee N",
            "annee_n1": "Annee N-1",
            "total_n": "Total N (EUR)",
            "total_n1": "Total N-1 (EUR)",
            "delta": "Delta (EUR)",
            "delta_pct": "Delta (%)",
        }
    )

    top_eur_affichage.to_excel(writer, sheet_name="Synthese", index=False, startrow=1)
    feuille = writer.book["Synthese"]
    feuille.cell(row=1, column=1, value="Top 20 hausses (EUR)")

    ligne_depart_pct = len(top_eur_affichage) + 4
    top_pct_affichage.to_excel(writer, sheet_name="Synthese", index=False, startrow=ligne_depart_pct)
    feuille.cell(row=ligne_depart_pct, column=1, value="Top 20 hausses (%)")

    _ajuster_largeurs(feuille)

    entete_eur_ligne = 2
    entete_pct_ligne = ligne_depart_pct + 1
    _formater_bloc_montants(feuille, entete_eur_ligne, len(top_eur_affichage))
    _formater_bloc_montants(feuille, entete_pct_ligne, len(top_pct_affichage))


def _formater_entete(feuille) -> None:
    """Appliquer un style homogène sur la ligne d'entête principale."""
    fond = PatternFill(fill_type="solid", fgColor="D9E1F2")
    gras = Font(bold=True)
    for cellule in feuille[1]:
        if cellule.value is not None:
            cellule.font = gras
            cellule.fill = fond


def _activer_filtres_et_gel(feuille) -> None:
    """Activer les filtres automatiques et geler la première ligne."""
    if feuille.max_row >= 1 and feuille.max_column >= 1:
        feuille.auto_filter.ref = feuille.dimensions
        feuille.freeze_panes = "A2"


def _ajuster_largeurs(feuille) -> None:
    """Ajuster les largeurs de colonnes pour une meilleure lisibilité."""
    for colonne in feuille.columns:
        longueur_max = 0
        lettre = colonne[0].column_letter
        for cellule in colonne:
            if cellule.value is not None:
                longueur_max = max(longueur_max, len(str(cellule.value)))
        feuille.column_dimensions[lettre].width = min(max(longueur_max + 2, 12), 48)


def _formater_colonnes_montant(feuille, indices_colonnes: list[int], number_format: str) -> None:
    """Appliquer un format numérique sur plusieurs colonnes (index Excel 1-based)."""
    for index_col in indices_colonnes:
        for ligne in range(2, feuille.max_row + 1):
            cellule = feuille.cell(row=ligne, column=index_col)
            if isinstance(cellule.value, (int, float)):
                cellule.number_format = number_format


def _formater_colonne_pourcentage(feuille, index_colonne: int) -> None:
    """Appliquer un format pourcentage en convertissant les valeurs en ratio."""
    for ligne in range(2, feuille.max_row + 1):
        cellule = feuille.cell(row=ligne, column=index_colonne)
        if isinstance(cellule.value, (int, float)):
            cellule.value = cellule.value / 100
            cellule.number_format = _FORMAT_PCT


def _formater_bloc_montants(feuille, ligne_entete: int, nb_lignes: int) -> None:
    """Formater un bloc tabulaire du feuille Synthese en EUR et pourcentage."""
    if nb_lignes == 0:
        return

    noms_colonnes = [feuille.cell(row=ligne_entete, column=i).value for i in range(1, feuille.max_column + 1)]

    for idx, nom in enumerate(noms_colonnes, start=1):
        if nom is None:
            continue
        if "(EUR)" in str(nom):
            for ligne in range(ligne_entete + 1, ligne_entete + nb_lignes + 1):
                cellule = feuille.cell(row=ligne, column=idx)
                if isinstance(cellule.value, (int, float)):
                    cellule.number_format = _FORMAT_EURO
        if "(%)" in str(nom):
            for ligne in range(ligne_entete + 1, ligne_entete + nb_lignes + 1):
                cellule = feuille.cell(row=ligne, column=idx)
                if isinstance(cellule.value, (int, float)):
                    cellule.value = cellule.value / 100
                    cellule.number_format = _FORMAT_PCT
