"""Tests du module xlsx.generateur."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from compta_copro.xlsx.generateur import generer_classeur_excel


def test_generer_classeur_excel_cree_un_fichier(tmp_path: Path) -> None:
    """Verifier que l'export XLSX cree le fichier et les onglets attendus."""
    df = pd.DataFrame(
        {
            "date": pd.to_datetime(["2024-03-01", "2025-03-01", "2025-04-01"]),
            "cle_repartition": ["001", "001", "001"],
            "type_charge": ["100 - ENTRETIEN", "100 - ENTRETIEN", "110 - VERTS"],
            "libelle": ["A", "B", "C"],
            "a_repartir": [100.0, 120.0, 80.0],
            "tva": [10.0, 12.0, 8.0],
            "recuperable": [100.0, 120.0, 80.0],
            "annee": [2024, 2025, 2025],
        }
    )

    sortie = tmp_path / "rapport.xlsx"
    chemin = generer_classeur_excel(df, sortie)

    assert chemin.exists()

    wb = openpyxl.load_workbook(chemin)
    try:
        assert "2024" in wb.sheetnames
        assert "2025" in wb.sheetnames
        assert "Comparaison" in wb.sheetnames
        assert "Synthese" in wb.sheetnames
    finally:
        wb.close()


def test_generer_classeur_excel_annee_courante(tmp_path: Path) -> None:
    """Verifier le mode restreint a l'annee courante."""
    df = pd.DataFrame(
        {
            "date": pd.to_datetime(["2023-03-01", "2024-03-01", "2025-03-01"]),
            "cle_repartition": ["001", "001", "001"],
            "type_charge": ["100 - ENTRETIEN", "100 - ENTRETIEN", "100 - ENTRETIEN"],
            "libelle": ["A", "B", "C"],
            "a_repartir": [90.0, 100.0, 120.0],
            "tva": [9.0, 10.0, 12.0],
            "recuperable": [90.0, 100.0, 120.0],
            "annee": [2023, 2024, 2025],
        }
    )

    sortie = tmp_path / "rapport_courant.xlsx"
    chemin = generer_classeur_excel(df, sortie, annee_cible=2025)

    wb = openpyxl.load_workbook(chemin)
    try:
        assert "2025" in wb.sheetnames
        assert "2024" not in wb.sheetnames
        assert "2023" not in wb.sheetnames
        assert "Comparaison" in wb.sheetnames
        assert "Synthese" in wb.sheetnames

        ws_comp = wb["Comparaison"]
        valeurs_annee_n = [ws_comp.cell(row=i, column=2).value for i in range(2, ws_comp.max_row + 1)]
        assert valeurs_annee_n
        assert all(v == 2025 for v in valeurs_annee_n if v is not None)
    finally:
        wb.close()
